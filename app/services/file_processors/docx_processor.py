import io
from docx import Document
from typing import Dict, Any
from app.services.file_processors.base_processor import BaseFileProcessor
from app.core.logging_config import logger


class DOCXProcessor(BaseFileProcessor):
    """Extract text from Word documents"""

    async def extract_text(self, file_path: str, file_content: bytes) -> Dict[str, Any]:
        """Extract text from DOCX with paragraph structure"""

        try:
            content = []
            warnings = []

            doc = Document(io.BytesIO(file_content))

            # Extract paragraphs
            paragraph_texts = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    paragraph_texts.append(text)

            # Extract tables
            table_texts = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        table_texts.append(row_text)

            # Combine paragraphs into sections
            if paragraph_texts:
                full_text = "\n\n".join(paragraph_texts)
                cleaned_text = self.clean_text(full_text)

                content.append(
                    {
                        "text": cleaned_text,
                        "metadata": {
                            "source": "docx",
                            "type": "paragraphs",
                            "paragraph_count": len(paragraph_texts),
                        },
                    }
                )

            # Add table data separately
            if table_texts:
                table_text = "\n".join(table_texts)
                cleaned_table = self.clean_text(table_text)

                content.append(
                    {
                        "text": cleaned_table,
                        "metadata": {
                            "source": "docx",
                            "type": "tables",
                            "table_count": len(doc.tables),
                        },
                    }
                )

            if not content:
                raise ValueError("No text could be extracted from Word document")

            return {
                "content": content,
                "metadata": {
                    "extraction_method": "python-docx",
                    "total_paragraphs": len(paragraph_texts),
                    "total_tables": len(doc.tables),
                    "warnings": warnings if warnings else None,
                },
            }

        except Exception as e:
            logger.error(f"DOCX processing failed: {str(e)}")
            raise ValueError(f"Failed to process Word document: {str(e)}")
